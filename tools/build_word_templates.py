from pathlib import Path
import random

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "plantillas_power_automate"
LOGO = ROOT / "assets" / "csc-logo.png"


BLUE = "004F8B"
LIGHT_BLUE = "EAF2FB"
LIGHT_GRAY = "D9E2EC"
TEXT = RGBColor(31, 41, 55)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(int(width_inches * 1440)))
    width.set(qn("w:type"), "dxa")


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def add_text(paragraph, text, bold=False, size=9, color=TEXT):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return run


def content_control_run(tag, placeholder=None):
    placeholder = placeholder or f"<<{tag}>>"
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")

    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), tag)
    sdt_pr.append(alias)

    tag_el = OxmlElement("w:tag")
    tag_el.set(qn("w:val"), tag)
    sdt_pr.append(tag_el)

    sdt_id = OxmlElement("w:id")
    sdt_id.set(qn("w:val"), str(random.randint(100000, 999999999)))
    sdt_pr.append(sdt_id)

    text_type = OxmlElement("w:text")
    sdt_pr.append(text_type)

    sdt_content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "444444")
    r_pr.append(color)
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = placeholder
    run.append(text)
    sdt_content.append(run)

    sdt.append(sdt_pr)
    sdt.append(sdt_content)
    return sdt


def add_content_control(paragraph, tag, placeholder=None):
    paragraph._p.append(content_control_run(tag, placeholder))


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.size = Pt(12 if level == 1 else 10.5)
    return paragraph


def add_field_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for label, tag in rows:
        cells = table.add_row().cells
        set_cell_width(cells[0], 2.35)
        set_cell_width(cells[1], 4.35)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cells[0], LIGHT_BLUE)
        set_cell_border(cells[0])
        set_cell_border(cells[1])

        left = cells[0].paragraphs[0]
        clear_paragraph(left)
        add_text(left, label, bold=True, size=8.8)

        right = cells[1].paragraphs[0]
        clear_paragraph(right)
        add_content_control(right, tag)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_checklist_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for label, tag in rows:
        cells = table.add_row().cells
        set_cell_width(cells[0], 4.75)
        set_cell_width(cells[1], 1.95)
        set_cell_border(cells[0])
        set_cell_border(cells[1])
        set_cell_shading(cells[1], "F7FAFC")

        left = cells[0].paragraphs[0]
        clear_paragraph(left)
        add_text(left, label, size=8.8)

        right = cells[1].paragraphs[0]
        clear_paragraph(right)
        right.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_content_control(right, tag)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def setup_doc(title):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)

    header_table = doc.add_table(rows=1, cols=2)
    header_table.autofit = False
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_width(header_table.cell(0, 0), 1.2)
    set_cell_width(header_table.cell(0, 1), 5.5)
    if LOGO.exists():
        p_logo = header_table.cell(0, 0).paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_logo.add_run().add_picture(str(LOGO), width=Inches(0.9))
    p_title = header_table.cell(0, 1).paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_title.add_run(title)
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_paragraph()
    return doc


COMMON_REQUEST_FIELDS = [
    ("Empresa solicitante", "empresa_solicitante"),
    ("Codigo de empresa", "empresa_codigo"),
    ("Departamento solicitante", "departamento_solicitante"),
    ("Solicitado por", "solicitado_por"),
    ("Correo del solicitante", "correo_solicitante"),
    ("Fecha de solicitud", "fecha_solicitud"),
    ("Tipo de solicitud", "tipo_solicitud"),
]


GENERAL_SUPPLIER_FIELDS = [
    ("Nombre comercial", "nombre_comercial"),
    ("Razon social", "razon_social"),
    ("NIT / Identificador fiscal", "nit_identificador_fiscal"),
    ("Tipo de persona", "tipo_persona_proveedor"),
    ("Direccion fiscal", "direccion_fiscal"),
    ("Nombre de contacto", "nombre_contacto"),
    ("Puesto del contacto", "puesto_contacto"),
    ("Telefono del contacto", "telefono_contacto"),
    ("Correo electronico", "correo_electronico"),
    ("Pagina web", "pagina_web"),
    ("Grupo", "grupo"),
    ("Moneda", "moneda_proveedor"),
    ("Dias de credito", "dias_credito"),
    ("Actividad economica / Giro del negocio", "actividad_economica"),
    ("Fecha de inicio de relacion comercial", "fecha_inicio_relacion"),
]


GENERAL_CLIENT_FIELDS = [
    ("Nombre comercial", "nombre_comercial"),
    ("Razon social", "razon_social"),
    ("NIT / Identificador fiscal", "nit_identificador_fiscal"),
    ("Tipo de persona", "tipo_persona_cliente"),
    ("Direccion fiscal", "direccion_fiscal"),
    ("Nombre contacto comercial", "nombre_contacto_comercial"),
    ("Puesto contacto comercial", "puesto_contacto_comercial"),
    ("Telefono contacto comercial", "telefono_contacto_comercial"),
    ("Correo contacto comercial", "correo_contacto_comercial"),
    ("Nombre contacto contabilidad", "nombre_contacto_contabilidad"),
    ("Puesto contacto contabilidad", "puesto_contacto_contabilidad"),
    ("Telefono contacto contabilidad", "telefono_contacto_contabilidad"),
    ("Correo contacto contabilidad", "correo_contacto_contabilidad"),
    ("Pagina web", "pagina_web"),
    ("Grupo", "grupo"),
    ("Moneda", "moneda_cliente"),
    ("Dias de credito", "dias_credito"),
    ("Actividad economica / Giro del negocio", "actividad_economica"),
    ("Fecha de inicio de relacion comercial", "fecha_inicio_relacion"),
]


TAX_SUPPLIER_FIELDS = [
    ("ISR sobre utilidades", "proveedor_isr_utilidades"),
    ("ISR opcional simplificado", "proveedor_isr_simplificado"),
    ("Pequeno contribuyente", "proveedor_pequeno_contribuyente"),
    ("Exento", "proveedor_exento"),
    ("Es agente retenedor", "proveedor_agente_retenedor"),
    ("Aplica retencion de ISR local", "proveedor_retencion_isr_local"),
    ("Aplica retencion de ISR del exterior", "proveedor_retencion_isr_exterior"),
    ("Aplica retencion de IVA", "proveedor_retencion_iva"),
]


TAX_CLIENT_FIELDS = [
    ("ISR sobre utilidades", "cliente_isr_utilidades"),
    ("ISR opcional simplificado", "cliente_isr_simplificado"),
    ("Pequeno contribuyente", "cliente_pequeno_contribuyente"),
    ("Exento", "cliente_exento"),
]


BANK_FIELDS = [
    ("Tipo de proveedor bancario", "tipo_proveedor_bancario"),
    ("Banco nacional", "banco_local"),
    ("Cuenta nacional", "cuenta_local"),
    ("Tipo de cuenta nacional", "tipo_cuenta_local_proveedor"),
    ("Moneda cuenta nacional", "moneda_cuenta_local_proveedor"),
    ("Nombre cuenta nacional", "nombre_cuenta_local"),
    ("Correo notificacion pago nacional", "correo_notificacion_local"),
    ("Banco extranjero", "banco_extranjero"),
    ("Pais del banco", "pais_banco"),
    ("Direccion del banco", "direccion_banco"),
    ("Cuenta extranjera", "cuenta_extranjera"),
    ("SWIFT / BIC", "swift_bic"),
    ("Tipo de cuenta extranjera", "tipo_cuenta_extranjera"),
    ("Moneda cuenta extranjera", "moneda_cuenta_ext_proveedor"),
    ("Otra moneda", "moneda_otra"),
    ("Nombre cuenta extranjera", "nombre_cuenta_extranjera"),
    ("Correo notificacion pago extranjera", "correo_notificacion_extranjera"),
]


CLIENT_BANK_FIELDS = [
    ("Tipo de cliente bancario", "tipo_cliente_bancario"),
    ("Banco nacional", "banco_local"),
    ("Cuenta nacional", "cuenta_local"),
    ("Tipo de cuenta nacional", "tipo_cuenta_local_cliente"),
    ("Moneda cuenta nacional", "moneda_cuenta_local_cliente"),
    ("Nombre cuenta nacional", "nombre_cuenta_local"),
    ("Correo notificacion pago nacional", "correo_notificacion_local"),
    ("Banco extranjero", "banco_extranjero"),
    ("Pais del banco", "pais_banco"),
    ("Direccion del banco", "direccion_banco"),
    ("Cuenta extranjera", "cuenta_extranjera"),
    ("SWIFT / BIC", "swift_bic"),
    ("Tipo de cuenta extranjera", "tipo_cuenta_extranjera"),
    ("Moneda cuenta extranjera", "moneda_cuenta_ext_cliente"),
    ("Otra moneda", "moneda_otra"),
    ("Nombre cuenta extranjera", "nombre_cuenta_extranjera"),
    ("Correo notificacion pago extranjera", "correo_notificacion_extranjera"),
]


REFERENCE_FIELDS = [
    ("Referencia 1 - compania", "ref1_compania"),
    ("Referencia 1 - contacto", "ref1_contacto"),
    ("Referencia 1 - telefono", "ref1_telefono"),
    ("Referencia 1 - correo", "ref1_correo"),
    ("Referencia 2 - compania", "ref2_compania"),
    ("Referencia 2 - contacto", "ref2_contacto"),
    ("Referencia 2 - telefono", "ref2_telefono"),
    ("Referencia 2 - correo", "ref2_correo"),
    ("Referencia 3 - compania", "ref3_compania"),
    ("Referencia 3 - contacto", "ref3_contacto"),
    ("Referencia 3 - telefono", "ref3_telefono"),
    ("Referencia 3 - correo", "ref3_correo"),
]


DOC_FIELDS_SUPPLIER = [
    ("RTU actualizado / Identificador fiscal", "proveedor_doc_juridica_rtu"),
    ("Patente de comercio", "proveedor_doc_juridica_patente"),
    ("Identificacion representante legal", "proveedor_doc_juridica_identificacion"),
    ("Nombramiento vigente", "proveedor_doc_juridica_nombramiento"),
    ("Certificacion bancaria", "proveedor_doc_juridica_bancaria"),
    ("Formulario de debida diligencia", "proveedor_doc_juridica_diligencia"),
    ("Documentacion de beneficiario final", "proveedor_doc_juridica_beneficiario"),
]


DOC_FIELDS_CLIENT = [
    ("RTU actualizado / Identificador fiscal", "cliente_doc_juridica_rtu"),
    ("Patente de comercio", "cliente_doc_juridica_patente"),
    ("Identificacion representante legal", "cliente_doc_juridica_identificacion"),
    ("Nombramiento vigente", "cliente_doc_juridica_nombramiento"),
    ("Certificacion bancaria", "cliente_doc_juridica_bancaria"),
    ("Formulario de debida diligencia", "cliente_doc_juridica_diligencia"),
    ("Documentacion de beneficiario final", "cliente_doc_juridica_beneficiario"),
]


def build_supplier_template():
    doc = setup_doc("FORMULARIO CREACION DE PROVEEDOR")

    add_heading(doc, "Datos de solicitud")
    add_field_table(doc, COMMON_REQUEST_FIELDS)

    add_heading(doc, "Datos generales del proveedor")
    add_field_table(doc, GENERAL_SUPPLIER_FIELDS)

    add_heading(doc, "Informacion tributaria")
    add_field_table(doc, TAX_SUPPLIER_FIELDS)

    add_heading(doc, "Datos bancarios")
    add_field_table(doc, BANK_FIELDS)

    add_heading(doc, "Documentacion requerida")
    add_checklist_table(doc, DOC_FIELDS_SUPPLIER)

    add_heading(doc, "Referencias comerciales")
    add_field_table(doc, REFERENCE_FIELDS)

    add_heading(doc, "Confirmacion del solicitante")
    add_field_table(
        doc,
        [
            ("Nombre del solicitante", "nombre_confirmacion_solicitante"),
            ("Confirmacion del solicitante", "confirmacion_solicitante_proveedor"),
        ],
    )

    add_heading(doc, "Uso exclusivo para datos maestros")
    add_field_table(
        doc,
        [
            ("Codigo del proveedor", "codigo_proveedor"),
            ("Observaciones", "observaciones_datos_maestros"),
        ],
    )
    return doc


def build_client_template():
    doc = setup_doc("FORMULARIO CREACION DE CLIENTE")

    add_heading(doc, "Datos de solicitud")
    add_field_table(doc, COMMON_REQUEST_FIELDS)

    add_heading(doc, "Datos generales del cliente")
    add_field_table(doc, GENERAL_CLIENT_FIELDS)

    add_heading(doc, "Informacion tributaria")
    add_field_table(doc, TAX_CLIENT_FIELDS)

    add_heading(doc, "Datos bancarios")
    add_field_table(doc, CLIENT_BANK_FIELDS)

    add_heading(doc, "Documentacion requerida")
    add_checklist_table(doc, DOC_FIELDS_CLIENT)

    add_heading(doc, "Referencias comerciales")
    add_field_table(doc, REFERENCE_FIELDS)

    add_heading(doc, "Confirmacion del solicitante")
    add_field_table(
        doc,
        [
            ("Nombre del solicitante", "nombre_confirmacion_solicitante"),
            ("Confirmacion del solicitante", "confirmacion_solicitante_cliente"),
        ],
    )

    add_heading(doc, "Uso exclusivo para datos maestros")
    add_field_table(
        doc,
        [
            ("Codigo del cliente", "codigo_cliente"),
            ("Observaciones", "observaciones_datos_maestros"),
        ],
    )
    return doc


def main():
    OUT_DIR.mkdir(exist_ok=True)
    outputs = [
        (OUT_DIR / "Plantilla_Proveedor_PowerAutomate.docx", build_supplier_template()),
        (OUT_DIR / "Plantilla_Cliente_PowerAutomate.docx", build_client_template()),
    ]
    for output, doc in outputs:
        doc.save(output)
        print(output)


if __name__ == "__main__":
    main()
